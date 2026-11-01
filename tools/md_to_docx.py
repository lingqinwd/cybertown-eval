"""把仓库里的 Markdown 报告转成可提交的 DOCX。

支持：# / ## / ### 标题、段落、无序列表、有序列表、引用、围栏代码块、
      Markdown 表格、**粗体**、行内 `代码`、水平线。
中文排版：正文宋体 11 pt，标题黑体，代码 Consolas 9 pt，表格 Table Grid。

用法：
    python tools/md_to_docx.py docs/PRACTICE_B_REPORT.md -o docs/PRACTICE_B_REPORT.docx \
        --title "工程实践报告（B 自命题型）" --subtitle "赛博小镇 —— 多智能体叙事系统的可治理性"
"""

from __future__ import annotations

import argparse
import datetime
import pathlib
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BODY_FONT_CN = "宋体"
BODY_FONT_EN = "Times New Roman"
HEAD_FONT_CN = "黑体"
CODE_FONT = "Consolas"


def set_font(run, cn=BODY_FONT_CN, en=BODY_FONT_EN, size=11, bold=False, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if color:
        run.font.color.rgb = RGBColor(*color)


def init_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT_EN
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_CN)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(6)

    for level, size in ((1, 18), (2, 15), (3, 13)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = BODY_FONT_EN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x33, 0x52)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT_CN)
        st.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        st.paragraph_format.space_after = Pt(8)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_rich_text(par, text: str, size=11, cn=BODY_FONT_CN):
    """处理 **粗体** 与 `行内代码`。"""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            set_font(par.add_run(piece[2:-2]), cn=cn, size=size, bold=True)
        elif piece.startswith("`") and piece.endswith("`"):
            set_font(par.add_run(piece[1:-1]), cn=cn, en=CODE_FONT, size=size - 1)
        else:
            set_font(par.add_run(piece), cn=cn, size=size)


def cover(doc: Document, title: str, subtitle: str):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(title), cn=HEAD_FONT_CN, size=26, bold=True,
             color=(0x1F, 0x33, 0x52))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(subtitle), cn=HEAD_FONT_CN, size=14, color=(0x44, 0x4A, 0x55))

    for _ in range(3):
        doc.add_paragraph()

    info = [
        ("课程", "AI 科研训练 · 工程实践 B（自命题型）"),
        ("开源仓库", "https://github.com/lingqinwd/cybertown-eval"),
        ("云端平台", "AUP Learning Cloud（JupyterLab 容器）"),
        ("姓名 / 学号", "__________  /  __________"),
        ("提交日期", datetime.date.today().isoformat()),
    ]
    for k, v in info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(f"{k}："), size=11, bold=True)
        set_font(p.add_run(v), size=11)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(cols):
            text = row[ci] if ci < len(row) else ""
            par = cells[ci].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            if ri == 0:
                add_rich_text(par, text, size=10)
                for run in par.runs:
                    run.bold = True
            else:
                add_rich_text(par, text, size=10)
    doc.add_paragraph()


def add_code(doc: Document, lines: list[str]):
    for line in lines:
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Pt(12)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.05
        set_font(par.add_run(line if line.strip() else " "),
                 en=CODE_FONT, cn=CODE_FONT, size=9)
    doc.add_paragraph()


def convert(md_path: pathlib.Path, out_path: pathlib.Path,
            title: str, subtitle: str):
    text = md_path.read_text(encoding="utf-8")
    doc = Document()
    init_styles(doc)
    cover(doc, title, subtitle)

    lines = text.split("\n")
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            # 去掉 |---| 分隔行
            body = [r for r in table_buf
                    if not all(re.fullmatch(r":?-{2,}:?", c.strip() or "-") for c in r)]
            add_table(doc, body)
            table_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                add_code(doc, code_buf)
                code_buf, in_code = [], False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            continue
        else:
            flush_table()

        if not stripped:
            i += 1
            continue

        # 水平线
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 3)
            h = doc.add_heading(level=level)
            set_font(h.add_run(m.group(2)), cn=HEAD_FONT_CN,
                     size={1: 18, 2: 15, 3: 13}[level], bold=True,
                     color=(0x1F, 0x33, 0x52))
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.space_after = Pt(6)
            add_rich_text(par, stripped.lstrip("> ").strip(), size=10)
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x50, 0x5A, 0x66)
            i += 1
            continue

        # 列表
        m = re.match(r"^([-*+])\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            add_rich_text(par, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Number")
            par.paragraph_format.space_after = Pt(3)
            add_rich_text(par, m.group(2))
            i += 1
            continue

        # 普通段落
        par = doc.add_paragraph()
        par.paragraph_format.first_line_indent = Pt(0)
        add_rich_text(par, stripped)
        i += 1

    flush_table()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"[OK] {out_path}  ({out_path.stat().st_size // 1024} KB)")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("-o", "--out", default="")
    ap.add_argument("--title", default="工程实践报告")
    ap.add_argument("--subtitle", default="")
    args = ap.parse_args()

    src = pathlib.Path(args.source)
    out = pathlib.Path(args.out) if args.out else src.with_suffix(".docx")
    convert(src, out, args.title, args.subtitle)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
